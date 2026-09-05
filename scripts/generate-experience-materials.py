from pathlib import Path
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pptx import Presentation
from pptx.util import Inches, Pt
import json, csv

root = Path(__file__).resolve().parents[1] / "体验材料" / "线性代数基础班"
root.mkdir(parents=True, exist_ok=True)

(root / "线性代数基础班_学生名单.csv").write_text("姓名,邮箱\n张三,zhangsan@example.com\n李四,lisi@example.com\n王五,wangwu@example.com\n赵六,zhaoliu@example.com\n陈可,chenke@example.com\n", encoding="utf-8-sig")
(root / "第1章_矩阵基础_公式速查.txt").write_text("线性代数基础班｜第1章 矩阵基础\n\n1. 矩阵加法：同型矩阵对应元素相加。\n2. 矩阵乘法：A(m×n)B(n×p) 可乘，结果为 m×p。\n3. 单位矩阵 I：AI=IA=A。\n4. 转置：(AB)^T=B^T A^T。\n5. 二阶行列式：|a b; c d|=ad-bc。\n", encoding="utf-8")
(root / "第1章_矩阵基础_课堂活动题目.json").write_text(json.dumps({"title":"矩阵基础课堂活动","questions":[{"type":"choice","prompt":"若 A 为 2×3 矩阵，B 为 3×1 矩阵，则 AB 的阶数是？","options":["2×1","3×2","2×3","不能相乘"],"answer":"2×1","knowledgePoint":"矩阵乘法的阶数"},{"type":"short_answer","prompt":"用一句话说明单位矩阵在矩阵乘法中的作用。","answer":"单位矩阵不改变矩阵"}]}, ensure_ascii=False, indent=2), encoding="utf-8")
(root / "个性化作业题库.json").write_text(json.dumps({"course":"线性代数","chapter":"第1章 矩阵基础","questions":[{"id":"la_q1","stem":"下列哪一个是 2×2 矩阵？","options":["[[1,0],[0,1]]","[1,2,3]","[[1,2,3]]","1"],"answer":"[[1,0],[0,1]]","knowledgePoint":"矩阵的阶数","difficulty":"easy"},{"id":"la_q2","stem":"若 A 为 2×3、B 为 3×4，则 AB 的阶数为？","options":["2×4","3×3","4×2","不能相乘"],"answer":"2×4","knowledgePoint":"矩阵乘法的阶数","difficulty":"medium"},{"id":"la_q3","stem":"矩阵转置满足哪条公式？","options":["(AB)^T=A^TB^T","(AB)^T=B^TA^T","(AB)^T=AB","(AB)^T=-AB"],"answer":"(AB)^T=B^TA^T","knowledgePoint":"矩阵转置","difficulty":"medium"}]}, ensure_ascii=False, indent=2), encoding="utf-8")

doc = Document(); doc.add_heading("第1章 矩阵基础｜例题与作业", 0); doc.add_paragraph("线性代数基础班 · 教师示例资料")
for title, body in [("例题1：矩阵乘法阶数", "设 A 为 2×3 矩阵，B 为 3×4 矩阵。由于 A 的列数等于 B 的行数，AB 可以相乘，结果为 2×4 矩阵。"),("例题2：单位矩阵", "设 I 为三阶单位矩阵，则对任意三阶矩阵 A，有 AI=IA=A。单位矩阵在矩阵乘法中相当于数字 1。"),("课后作业", "1. 判断 A(3×2) 与 B(2×5) 是否可乘，并写出 AB 的阶数。\n2. 计算 [[1,2],[0,1]] 的转置。\n3. 解释为什么 (AB)^T=B^TA^T。")]:
    doc.add_heading(title, 1); doc.add_paragraph(body)
doc.save(root / "第1章_矩阵基础_例题.docx")

pdfmetrics.registerFont(TTFont("SimHei", "C:/Windows/Fonts/simhei.ttf"))
c = canvas.Canvas(str(root / "第1章_矩阵基础_讲义.pdf"), pagesize=A4); w, h = A4; y = h - 60
c.setFont("SimHei", 20); c.drawString(55, y, "第1章 矩阵基础"); y -= 35; c.setFont("SimHei", 11)
for line in ["线性代数基础班 · 讲义", "", "一、矩阵的阶数", "矩阵有 m 行 n 列时记作 m×n。判断两个矩阵能否相乘，关键是前者列数等于后者行数。", "", "二、矩阵乘法", "A(m×n)B(n×p) 的结果是 m×p。结果矩阵第 i,j 个元素等于 A 第 i 行与 B 第 j 列的内积。", "", "三、单位矩阵与转置", "单位矩阵 I 满足 AI=IA=A；转置满足 (AB)^T=B^T A^T。", "", "思考：若 A 是 2×3、B 是 3×1，AB 的阶数是什么？"]:
    c.drawString(55, y, line); y -= 20
c.save()

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
for title, body in [("第1章 矩阵基础", "线性代数基础班\n从阶数、乘法到转置"),("矩阵乘法先看阶数", "A(m×n)B(n×p) 可乘\n结果是 m×p\n示例：2×3 · 3×4 = 2×4"),("课堂活动", "若 A 为 2×3，B 为 3×1，AB 的阶数是？\nA. 2×1   B. 3×2   C. 2×3   D. 不能相乘"),("课后迁移", "把一道题的答案、过程和一个仍然困惑的地方提交给老师。")]:
    slide = prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text = title; slide.placeholders[1].text = body
    for para in slide.placeholders[1].text_frame.paragraphs: para.font.size = Pt(24)
prs.save(root / "第1章_矩阵基础_课堂课件.pptx")

(root / "矩阵基础作业_学生答案示例.docx").write_text("", encoding="utf-8")
doc = Document(); doc.add_heading("矩阵基础作业｜学生答案示例", 0); doc.add_paragraph("学生：张三"); doc.add_heading("第1题", 1); doc.add_paragraph("A(3×2) 与 B(2×5) 可以相乘，AB 的阶数是 3×5。"); doc.add_heading("第2题", 1); doc.add_paragraph("[[1,2],[0,1]] 的转置为 [[1,0],[2,1]]。"); doc.add_heading("我的疑问", 1); doc.add_paragraph("我想进一步理解转置后乘法顺序为什么会反过来。")
doc.save(root / "矩阵基础作业_学生答案示例.docx")
print(root)
